from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import json
import zipfile

from openpyxl import load_workbook


class RepairFinanceError(ValueError):
    pass


DEFAULT_POLICY = {
    "eligible_statuses": ["维修完成"],
    "employee_rules": [
        {"name": "李蒙悦", "base_salary": 3500, "commission_rate": 0.18, "subsidy": 51.1, "deduction": 700},
        {"name": "贺祥军", "base_salary": 3200, "commission_rate": 0.18, "subsidy": 352.34, "deduction": 0},
        {"name": "宋嘉庆", "base_salary": 1800, "commission_rate": 0, "subsidy": 0, "deduction": 0},
        {"name": "徐越", "base_salary": 5200, "commission_rate": 0, "subsidy": 0, "deduction": 0},
    ],
    "sales_commission_rate": 0.05,
    "partner_name": "鲁明涛",
    "partner_share_rate": 0.5,
    "operating_expense": 0,
}

REQUIRED_HEADERS = {"ID", "时间", "工程师", "报价", "成本"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


def _money(value: Decimal) -> float:
    return float(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _decimal(value: object, label: str) -> Decimal:
    if value is None or value == "":
        raise RepairFinanceError(f"{label}为空")
    if isinstance(value, bool):
        raise RepairFinanceError(f"{label}不是数字")
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise RepairFinanceError(f"{label}不是有效数字") from exc
    if not result.is_finite():
        raise RepairFinanceError(f"{label}不是有限数字")
    return result


def _rate(value: object, label: str) -> Decimal:
    result = _decimal(value, label)
    if result < 0 or result > 1:
        raise RepairFinanceError(f"{label}必须在0到1之间")
    return result


def normalize_policy(raw: dict | None) -> dict:
    source = {**DEFAULT_POLICY, **(raw or {})}
    statuses = [str(value).strip() for value in source.get("eligible_statuses", []) if str(value).strip()]
    if not statuses:
        raise RepairFinanceError("至少需要一个纳入结算的订单状态")
    employee_rules: list[dict] = []
    names: set[str] = set()
    for index, item in enumerate(source.get("employee_rules", []), start=1):
        name = str((item or {}).get("name") or "").strip()
        if not name:
            raise RepairFinanceError(f"第{index}条工资规则缺少姓名")
        if name in names:
            raise RepairFinanceError(f"工资规则中的员工重复：{name}")
        names.add(name)
        base = _decimal((item or {}).get("base_salary", 0), f"{name}基础工资")
        subsidy = _decimal((item or {}).get("subsidy", 0), f"{name}其他补贴")
        deduction = _decimal((item or {}).get("deduction", 0), f"{name}其他扣款")
        if min(base, subsidy, deduction) < 0:
            raise RepairFinanceError(f"{name}的工资、补贴和扣款不得为负数")
        employee_rules.append({
            "name": name,
            "base_salary": base,
            "commission_rate": _rate((item or {}).get("commission_rate", 0), f"{name}提成比例"),
            "subsidy": subsidy,
            "deduction": deduction,
        })
    partner_name = str(source.get("partner_name") or "").strip()
    operating_expense = _decimal(source.get("operating_expense", 0), "其他经营支出")
    if operating_expense < 0:
        raise RepairFinanceError("其他经营支出不得为负数")
    return {
        "eligible_statuses": statuses,
        "employee_rules": employee_rules,
        "sales_commission_rate": _rate(source.get("sales_commission_rate", 0), "销售提成比例"),
        "partner_name": partner_name,
        "partner_share_rate": _rate(source.get("partner_share_rate", 0), "合作分润比例"),
        "operating_expense": operating_expense,
    }


def policy_for_storage(raw: dict | None) -> dict:
    policy = normalize_policy(raw)
    return {
        "eligible_statuses": policy["eligible_statuses"],
        "employee_rules": [
            {
                "name": item["name"],
                "base_salary": _money(item["base_salary"]),
                "commission_rate": float(item["commission_rate"]),
                "subsidy": _money(item["subsidy"]),
                "deduction": _money(item["deduction"]),
            }
            for item in policy["employee_rules"]
        ],
        "sales_commission_rate": float(policy["sales_commission_rate"]),
        "partner_name": policy["partner_name"],
        "partner_share_rate": float(policy["partner_share_rate"]),
        "operating_expense": _money(policy["operating_expense"]),
    }


def _validate_xlsx(content: bytes) -> None:
    if not content:
        raise RepairFinanceError("上传文件为空")
    if len(content) > MAX_UPLOAD_BYTES:
        raise RepairFinanceError("文件超过10MB")
    stream = BytesIO(content)
    if not zipfile.is_zipfile(stream):
        raise RepairFinanceError("请上传有效的 .xlsx 文件")
    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        if len(archive.infolist()) > 5000:
            raise RepairFinanceError("工作簿内部文件数量异常")
        if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise RepairFinanceError("工作簿解压后过大")


def _find_detail_sheet(workbook):
    candidates = []
    if "订单详情" in workbook.sheetnames:
        candidates.append(workbook["订单详情"])
    candidates.extend(sheet for sheet in workbook.worksheets if sheet not in candidates)
    for sheet in candidates:
        headers = {str(cell.value).strip() for cell in sheet[1] if cell.value not in (None, "")}
        if REQUIRED_HEADERS.issubset(headers):
            return sheet
    raise RepairFinanceError("未找到订单明细表，第一行至少需要 ID、时间、工程师、报价、成本")


def analyze_repair_billing_xlsx(content: bytes, filename: str, raw_policy: dict | None = None) -> dict:
    _validate_xlsx(content)
    policy = normalize_policy(raw_policy)
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True, keep_links=False)
    except Exception as exc:
        raise RepairFinanceError("工作簿读取失败") from exc
    try:
        sheet = _find_detail_sheet(workbook)
        headers: dict[str, int] = {}
        for column, cell in enumerate(sheet[1], start=1):
            header = str(cell.value).strip() if cell.value not in (None, "") else ""
            if header and header not in headers:
                headers[header] = column

        def cell(row: tuple, name: str):
            column = headers.get(name)
            return row[column - 1].value if column and column <= len(row) else None

        eligible_statuses = set(policy["eligible_statuses"])
        partner_name = policy["partner_name"]
        employee_rules = {item["name"]: item for item in policy["employee_rules"]}
        seen_ids: set[str] = set()
        all_rows = 0
        accepted: list[dict] = []
        issues: list[dict] = []
        status_counts: Counter[str] = Counter()
        excluded_status_counts: Counter[str] = Counter()

        def issue(severity: str, code: str, row_number: int, order_id: str, message: str) -> None:
            issues.append({"severity": severity, "code": code, "row": row_number, "order_id": order_id, "message": message})

        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            raw_id = cell(row, "ID")
            if raw_id in (None, ""):
                if any(value.value not in (None, "") for value in row):
                    issue("error", "missing_id", row_number, "", "缺少订单ID，该行未纳入计算")
                continue
            order_id = str(raw_id).strip()
            if order_id == "汇总":
                continue
            all_rows += 1
            if order_id in seen_ids:
                issue("error", "duplicate_id", row_number, order_id, "订单ID重复，后出现的记录已排除")
                continue
            seen_ids.add(order_id)
            status = str(cell(row, "备注") or "").strip() or "（空白）"
            status_counts[status] += 1
            if status not in eligible_statuses:
                excluded_status_counts[status] += 1
                continue
            try:
                quote = _decimal(cell(row, "报价"), "报价")
                cost = _decimal(cell(row, "成本"), "成本")
            except RepairFinanceError as exc:
                issue("error", "invalid_amount", row_number, order_id, f"{exc}，该订单未纳入计算")
                continue
            if quote <= 0:
                issue("error", "non_positive_quote", row_number, order_id, "报价必须大于0，该订单未纳入计算")
                continue
            if cost < 0:
                issue("error", "negative_cost", row_number, order_id, "成本不得为负数，该订单未纳入计算")
                continue
            parts = [str(cell(row, key) or "").strip() for key in ("配件1", "配件2", "配件3")]
            part_source = str(cell(row, "配件来源") or "").strip()
            if cost == 0 and (any(parts) or (part_source and part_source != "无")):
                issue("error", "missing_part_cost", row_number, order_id, "记录了配件或配件来源，但成本为0，该订单未纳入计算")
                continue
            engineer = str(cell(row, "工程师") or "").strip()
            if not engineer:
                issue("error", "missing_engineer", row_number, order_id, "缺少工程师，该订单未纳入计算")
                continue
            if engineer != partner_name and engineer not in employee_rules:
                issue("error", "missing_employee_rule", row_number, order_id, f"工程师 {engineer} 没有工资规则，该订单未纳入计算")
                continue
            calculated_profit = quote - cost
            if calculated_profit < 0:
                issue("warning", "negative_order_profit", row_number, order_id, "成本高于报价，该订单产生负利润，已保留并提示复核")
            uploaded_profit = cell(row, "利润")
            if uploaded_profit not in (None, ""):
                try:
                    if abs(_decimal(uploaded_profit, "利润") - calculated_profit) > Decimal("0.01"):
                        issue("warning", "profit_recalculated", row_number, order_id, "上传利润与报价减成本不一致，已按报价减成本重算")
                except RepairFinanceError:
                    issue("warning", "profit_recalculated", row_number, order_id, "上传利润无效，已按报价减成本重算")
            accepted.append({
                "order_id": order_id,
                "row": row_number,
                "date": str(cell(row, "时间") or ""),
                "customer": str(cell(row, "顾客") or "").strip(),
                "engineer": engineer,
                "salesperson": str(cell(row, "销售姓名") or "").strip(),
                "quote": quote,
                "cost": cost,
                "profit": calculated_profit,
            })

        by_engineer: dict[str, list[dict]] = defaultdict(list)
        by_salesperson: dict[str, list[dict]] = defaultdict(list)
        for order in accepted:
            by_engineer[order["engineer"]].append(order)
            if order["salesperson"]:
                by_salesperson[order["salesperson"]].append(order)

        payroll = []
        labor_expense = Decimal("0")
        cash_payroll = Decimal("0")
        for name, rule in employee_rules.items():
            orders = by_engineer.get(name, [])
            order_profit = sum((item["profit"] for item in orders), Decimal("0"))
            commission = order_profit * rule["commission_rate"]
            gross_pay = rule["base_salary"] + commission + rule["subsidy"]
            net_pay = gross_pay - rule["deduction"]
            if net_pay < 0:
                issues.append({"severity": "error", "code": "negative_net_pay", "row": 0, "order_id": "", "message": f"{name}的扣款超过应发工资"})
            labor_expense += gross_pay
            cash_payroll += max(net_pay, Decimal("0"))
            payroll.append({
                "name": name,
                "order_count": len(orders),
                "order_profit": _money(order_profit),
                "base_salary": _money(rule["base_salary"]),
                "commission_rate": float(rule["commission_rate"]),
                "commission": _money(commission),
                "subsidy": _money(rule["subsidy"]),
                "deduction": _money(rule["deduction"]),
                "gross_pay": _money(gross_pay),
                "net_pay": _money(max(net_pay, Decimal("0"))),
            })

        sales = []
        sales_commission = Decimal("0")
        for name, orders in sorted(by_salesperson.items()):
            order_profit = sum((item["profit"] for item in orders), Decimal("0"))
            commission = order_profit * policy["sales_commission_rate"]
            sales_commission += commission
            sales.append({
                "name": name,
                "order_count": len(orders),
                "order_profit": _money(order_profit),
                "commission_rate": float(policy["sales_commission_rate"]),
                "commission": _money(commission),
            })

        revenue = sum((item["quote"] for item in accepted), Decimal("0"))
        direct_cost = sum((item["cost"] for item in accepted), Decimal("0"))
        gross_profit = revenue - direct_cost
        profit_before_share = gross_profit - labor_expense - sales_commission - policy["operating_expense"]
        share_base = max(profit_before_share, Decimal("0"))
        partner_share = share_base * policy["partner_share_rate"] if partner_name else Decimal("0")
        retained_profit = profit_before_share - partner_share
        if profit_before_share < 0:
            issues.append({"severity": "warning", "code": "negative_distributable_profit", "row": 0, "order_id": "", "message": "扣除人工、销售提成和经营支出后为亏损，合作分润按0计算"})

        error_count = sum(item["severity"] == "error" for item in issues)
        warning_count = sum(item["severity"] == "warning" for item in issues)
        return {
            "source": {
                "filename": filename,
                "sheet": sheet.title,
                "total_rows": all_rows,
                "accepted_orders": len(accepted),
                "excluded_orders": all_rows - len(accepted),
            },
            "policy": {
                "eligible_statuses": policy["eligible_statuses"],
                "sales_commission_rate": float(policy["sales_commission_rate"]),
                "partner_name": partner_name,
                "partner_share_rate": float(policy["partner_share_rate"]),
                "operating_expense": _money(policy["operating_expense"]),
            },
            "summary": {
                "order_count": len(accepted),
                "revenue": _money(revenue),
                "direct_cost": _money(direct_cost),
                "gross_profit": _money(gross_profit),
                "gross_margin": float((gross_profit / revenue).quantize(Decimal("0.0001"))) if revenue else 0,
                "labor_expense": _money(labor_expense),
                "cash_payroll": _money(cash_payroll),
                "sales_commission": _money(sales_commission),
                "operating_expense": _money(policy["operating_expense"]),
                "profit_before_share": _money(profit_before_share),
                "partner_share": _money(partner_share),
                "retained_profit": _money(retained_profit),
            },
            "payroll": payroll,
            "sales": sales,
            "status_summary": [{"status": key, "count": value, "included": key in eligible_statuses} for key, value in status_counts.most_common()],
            "excluded_status_summary": [{"status": key, "count": value} for key, value in excluded_status_counts.most_common()],
            "issues": issues,
            "error_count": error_count,
            "warning_count": warning_count,
            "can_finalize": error_count == 0,
        }
    finally:
        workbook.close()


def policy_from_json(value: str) -> dict:
    if not value.strip():
        return DEFAULT_POLICY
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as exc:
        raise RepairFinanceError("结算规则不是有效JSON") from exc
    if not isinstance(parsed, dict):
        raise RepairFinanceError("结算规则必须是对象")
    return parsed
